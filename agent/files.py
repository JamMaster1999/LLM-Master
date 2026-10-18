"""The file tools: read, write and edit copy pi's; deliver_file hands a file to the user."""
import asyncio
import base64
import codecs
import json
import mimetypes
import posixpath

from modal.exception import SandboxFilesystemError, SandboxFilesystemNotFoundError
from modal.types import FileType

from agent.result import Result, ToolError, UserFile
from agent.sandbox import Sandbox, resolve

MAX_LINES, MAX_BYTES = 2000, 50 * 1024  # pi's cut for long text
MAX_PAGES = 10
SHEET_BYTES = 25 * 1024  # a spreadsheet up to this much CSV comes back whole
MAX_FILE = 50 * 1024 * 1024  # the most any tool pulls out of the sandbox at once
RENDER_TIMEOUT = 120

IMAGES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic", ".heif"}
PAGED = {".pdf", ".docx", ".doc", ".odt", ".rtf", ".pptx", ".ppt", ".odp"}
SHEETS = {".xlsx", ".xlsm", ".xls", ".ods", ".csv", ".tsv"}
MEDIA = {".mp3", ".wav", ".m4a", ".aac", ".ogg", ".oga", ".flac", ".opus", ".mp4", ".m4v", ".mov", ".webm",
         ".mkv", ".avi"}

READ = {"type": "function", "function": {
    "name": "read",
    "description": "Read the contents of a file. Supports text files, images (jpg, png, gif, webp, bmp, tiff, heic), "
                   "PDFs and office documents (docx, pptx and others, shown as page images), and spreadsheets "
                   "(xlsx, xls, ods, csv, tsv, as CSV). Images are sent as attachments. For text files, output is "
                   "truncated to 2000 lines or 50KB (whichever is hit first). Use offset/limit for large files. When "
                   "you need the full file, continue with offset until complete. For PDFs and documents, offset and "
                   "limit count pages, at most 10 per call.",
    "parameters": {"type": "object", "properties": {
        "path": {"type": "string", "description": "Path to the file to read (relative or absolute)"},
        "offset": {"type": "integer", "description": "Line number to start reading from (1-indexed); for PDFs and documents, the first page"},
        "limit": {"type": "integer", "description": "Maximum number of lines to read; for PDFs and documents, pages (at most 10)"},
    }, "required": ["path"]}}}

WRITE = {"type": "function", "function": {
    "name": "write",
    "description": "Write content to a file. Creates the file if it doesn't exist, overwrites if it does. "
                   "Automatically creates parent directories.",
    "parameters": {"type": "object", "properties": {
        "path": {"type": "string", "description": "Path to the file to write (relative or absolute)"},
        "content": {"type": "string", "description": "Content to write to the file"},
    }, "required": ["path", "content"]}}}

EDIT = {"type": "function", "function": {
    "name": "edit",
    "description": "Edit a single file using exact text replacement. Every edits[].oldText must match a unique, "
                   "non-overlapping region of the original file. If two changes affect the same block or nearby "
                   "lines, merge them into one edit instead of emitting overlapping edits. Do not include large "
                   "unchanged regions just to connect distant changes.",
    "parameters": {"type": "object", "properties": {
        "path": {"type": "string", "description": "Path to the file to edit (relative or absolute)"},
        "edits": {"type": "array", "description": "One or more targeted replacements. Each edit is matched against "
                  "the original file, not incrementally. Do not include overlapping or nested edits. If two changes "
                  "touch the same block or nearby lines, merge them into one edit instead.",
                  "items": {"type": "object", "properties": {
                      "oldText": {"type": "string", "description": "Exact text for one targeted replacement. It must "
                                  "be unique in the original file and must not overlap with any other "
                                  "edits[].oldText in the same call."},
                      "newText": {"type": "string", "description": "Replacement text for this targeted edit."},
                  }, "required": ["oldText", "newText"]}},
    }, "required": ["path", "edits"]}}}

DELIVER = {"type": "function", "function": {
    "name": "deliver_file",
    "description": "Send a file from the sandbox to the user, who can open or download it. Use it for every file "
                   "the user asked for, once it is finished.",
    "parameters": {"type": "object", "properties": {
        "path": {"type": "string", "description": "Path to the file to send (relative or absolute)"},
    }, "required": ["path"]}}}


async def read(sandbox: Sandbox, path: str, offset: int | None = None, limit: int | None = None) -> Result:
    extension = posixpath.splitext(path)[1].lower()
    if extension in MEDIA:
        raise ToolError(f"read can't return audio or video. Use bash (ffprobe, ffmpeg) to work with {path}.")
    await _stat(sandbox, path)
    if extension in IMAGES:
        mime, image = (await _run(sandbox, path, _IMAGE, resolve(path))).split(" ", 1)
        return Result([{"type": "text", "text": f"Read image file [{mime}]"},
                       {"type": "image", "data": _base64(await sandbox.read(image)), "mime": mime}])
    if extension in PAGED:
        return await _pages(sandbox, path, offset or 1, min(limit or MAX_PAGES, MAX_PAGES))
    if extension in SHEETS:
        return Result(await _run(sandbox, path, _SHEET, resolve(path), str(SHEET_BYTES)))
    data = await _load(sandbox, path)
    if not _utf8(data[:8192]):
        raise ToolError(f"{path} is a binary file or not UTF-8 text. Use bash to inspect it.")
    return Result(page(data.decode(errors="replace"), path, offset, limit))


async def write(sandbox: Sandbox, path: str, content: str) -> str:
    try:
        await sandbox.write(path, content.encode())
    except SandboxFilesystemError as e:
        raise ToolError(f"Could not write {path}: {e}") from None
    return f"Successfully wrote to {path}"


async def edit(sandbox: Sandbox, path: str, edits: list[dict] | dict | str) -> str:
    if isinstance(edits, str):  # pi: some models send the array as a JSON string
        try:
            edits = json.loads(edits)
        except json.JSONDecodeError:
            raise ToolError("Edit tool input is invalid. edits must be an array of {oldText, newText}.") from None
    if isinstance(edits, dict):  # and some send one edit without the array
        edits = [edits]
    if not isinstance(edits, list) or not all(
            isinstance(e, dict) and isinstance(e.get("oldText"), str) and isinstance(e.get("newText"), str) for e in edits):
        raise ToolError("Edit tool input is invalid. edits must be an array of {oldText, newText}.")
    try:
        text = (await _load(sandbox, path)).decode()
    except UnicodeDecodeError:
        raise ToolError(f"{path} is not UTF-8 text. Use bash to change it.") from None
    await write(sandbox, path, apply_edits(text, edits, path))
    return f"Successfully replaced {len(edits)} block(s) in {path}."


async def deliver_file(sandbox: Sandbox, path: str) -> Result:
    data = await _load(sandbox, path)
    name = posixpath.basename(path)
    return Result(f"Sent {name} to the user.",
                  files=[UserFile(name, mimetypes.guess_type(name)[0] or "application/octet-stream", data)])


def page(text: str, path: str, offset: int | None, limit: int | None) -> str:
    """pi's read window: from `offset`, at most `limit` lines, cut at 2000 lines or 50 KB, with a
    notice that says how to read on."""
    lines = text.split("\n")
    if text.endswith("\n"):
        lines.pop()
    total = len(lines)
    start = (offset or 1) - 1
    if offset and start >= total:
        raise ToolError(f"Offset {offset} is beyond end of file ({total} lines total)")
    end = min(total, start + limit) if limit else total
    shown, size = [], 0
    for line in lines[start:end]:
        size += len(line.encode()) + 1
        if len(shown) == MAX_LINES or size > MAX_BYTES:
            break
        shown.append(line)
    if not shown and start < end:
        line = start + 1
        return (f"[Line {line} is {_size(len(lines[start].encode()))}, exceeds {_size(MAX_BYTES)} limit. "
                f"Use bash: sed -n '{line}p' {path} | head -c {MAX_BYTES}]")
    last = start + len(shown)
    out = "\n".join(shown)
    if last < end:
        cap = "" if len(shown) == MAX_LINES else f" ({_size(MAX_BYTES)} limit)"
        out += f"\n\n[Showing lines {start + 1}-{last} of {total}{cap}. Use offset={last + 1} to continue.]"
    elif end < total:
        out += f"\n\n[{total - end} more lines in file. Use offset={end + 1} to continue.]"
    return out


def apply_edits(text: str, edits: list[dict], path: str) -> str:
    """pi's edit: every oldText matches the original exactly once, no two overlap, and the file keeps
    its BOM and line endings."""
    if not edits:
        raise ToolError("Edit tool input is invalid. edits must contain at least one replacement.")
    one = len(edits) == 1
    bom = "﻿" if text.startswith("﻿") else ""
    original = text.removeprefix(bom)
    crlf = "\r\n" in original
    original = original.replace("\r\n", "\n")
    spans = []
    for i, edit in enumerate(edits):
        old, new = edit["oldText"].replace("\r\n", "\n"), edit["newText"].replace("\r\n", "\n")
        name = "the text" if one else f"edits[{i}]"
        if not old:
            raise ToolError(f"oldText must not be empty in {path}." if one else f"edits[{i}].oldText must not be empty in {path}.")
        count = original.count(old)
        if count == 0:
            raise ToolError(f"Could not find the exact text in {path}. The old text must match exactly including all whitespace and newlines."
                            if one else f"Could not find {name} in {path}. The oldText must match exactly including all whitespace and newlines.")
        if count > 1:
            raise ToolError(f"Found {count} occurrences of the text in {path}. The text must be unique. Please provide more context to make it unique."
                            if one else f"Found {count} occurrences of {name} in {path}. Each oldText must be unique. Please provide more context to make it unique.")
        start = original.index(old)
        spans.append((start, start + len(old), new, i))
    spans.sort()
    for (_, end, _, a), (start, _, _, b) in zip(spans, spans[1:]):
        if start < end:
            raise ToolError(f"edits[{a}] and edits[{b}] overlap in {path}. Merge them into one edit or target disjoint regions.")
    changed = original
    for start, end, new, _ in reversed(spans):
        changed = changed[:start] + new + changed[end:]
    if changed == original:
        raise ToolError(f"No changes made to {path}. The replacement produced identical content. This might indicate an issue with special characters or the text not existing as expected."
                        if one else f"No changes made to {path}. The replacements produced identical content.")
    return bom + (changed.replace("\n", "\r\n") if crlf else changed)


async def _stat(sandbox: Sandbox, path: str) -> int:
    """The size of a regular file, or a ToolError that says what's wrong with the path."""
    try:
        info = await sandbox.stat(path)
    except SandboxFilesystemNotFoundError:
        raise ToolError(f"File not found: {path}") from None
    if info.type == FileType.DIRECTORY:
        raise ToolError(f"{path} is a directory. Use bash to list it.")
    if info.type != FileType.FILE:
        raise ToolError(f"{path} is not a regular file.")
    return info.size


async def _load(sandbox: Sandbox, path: str) -> bytes:
    size = await _stat(sandbox, path)
    if size > MAX_FILE:
        raise ToolError(f"{path} is {_size(size)}, over the {_size(MAX_FILE)} limit. Use bash to work with it.")
    return await sandbox.read(path)


async def _pages(sandbox: Sandbox, path: str, first: int, count: int) -> Result:
    if first < 1 or count < 1:
        raise ToolError("offset and limit start at 1.")
    total, *pages = (await _run(sandbox, path, _RENDER, resolve(path), str(first), str(count))).split("\n")
    if first > int(total):
        raise ToolError(f"Offset {first} is beyond the last page of {path} ({total} pages)")
    last = first + len(pages) - 1
    header = f"Rendered {path} [pages {first}-{last} of {total}]"
    if last < int(total):
        header += f" Use offset={last + 1} to continue."
    images = await asyncio.gather(*(sandbox.read(page) for page in pages))
    return Result([{"type": "text", "text": header},
                   *({"type": "image", "data": _base64(image), "mime": "image/jpeg"} for image in images)])


async def _run(sandbox: Sandbox, path: str, script: str, *args: str) -> str:
    """Runs one of the scripts below in the sandbox; returns what it printed."""
    try:
        code, stdout, stderr = await sandbox.exec("python", "-c", script, *args, timeout=RENDER_TIMEOUT)
    except TimeoutError:
        raise ToolError(f"Reading {path} took over {RENDER_TIMEOUT} seconds.") from None
    if code:
        reason = stderr.decode(errors="replace").strip().splitlines() or ["unknown error"]
        raise ToolError(f"Could not read {path}: {reason[-1]}")
    return stdout.decode(errors="replace").strip()


def _utf8(head: bytes) -> bool:
    """Whether the start of a file reads as UTF-8 text; a character cut at the end doesn't count."""
    try:
        codecs.getincrementaldecoder("utf-8")().decode(head)
    except UnicodeDecodeError:
        return False
    return b"\0" not in head


def _base64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def _size(n: int) -> str:
    """Sizes the way pi writes them: 512B, 50.0KB, 1.2MB."""
    if n < 1024:
        return f"{n}B"
    if n < 1024 * 1024:
        return f"{n / 1024:.1f}KB"
    return f"{n / 1024 / 1024:.1f}MB"


# Scripts that run inside the sandbox, where Pillow, openpyxl, LibreOffice and poppler live.

_IMAGE = r"""
import sys, tempfile
from PIL import Image
import pillow_heif
pillow_heif.register_heif_opener()
src = sys.argv[1]
image = Image.open(src)
if image.format in ("PNG", "JPEG", "GIF", "WEBP") and max(image.size) <= 1568:
    print(Image.MIME[image.format], src)
else:
    image.thumbnail((1568, 1568))
    alpha = image.mode in ("RGBA", "LA") or "transparency" in image.info
    out = tempfile.mkstemp(suffix=".png" if alpha else ".jpg")[1]
    image.convert("RGBA" if alpha else "RGB").save(out, quality=85)
    print("image/png" if alpha else "image/jpeg", out)
"""

_RENDER = r"""
import glob, subprocess, sys, tempfile
src, first, count = sys.argv[1], int(sys.argv[2]), int(sys.argv[3])
out = tempfile.mkdtemp()
def run(*args):
    done = subprocess.run(args, capture_output=True, text=True)
    if done.returncode:
        sys.exit(done.stderr.strip() or f"{args[0]} failed")
    return done.stdout
if not src.lower().endswith(".pdf"):  # a profile of its own, so parallel conversions don't fight over one
    run("soffice", f"-env:UserInstallation=file://{out}/profile", "--headless", "--convert-to", "pdf", "--outdir", out, src)
    src = glob.glob(f"{out}/*.pdf")[0]
total = int(next(line.split()[1] for line in run("pdfinfo", src).splitlines() if line.startswith("Pages:")))
if first <= total:
    run("pdftoppm", "-jpeg", "-jpegopt", "quality=85", "-scale-to", "1568", "-f", str(first),
        "-l", str(min(first + count - 1, total)), src, f"{out}/page")
print(total)
print("\n".join(sorted(glob.glob(f"{out}/page*.jpg"))))
"""

_SHEET = r"""
import csv, glob, io, os, subprocess, sys, tempfile, warnings
warnings.simplefilter("ignore")
src, limit = sys.argv[1], int(sys.argv[2])
name = os.path.basename(src)
if src.lower().endswith((".csv", ".tsv")):
    with open(src, newline="", encoding="utf-8", errors="replace") as f:
        sheets = {name: list(csv.reader(f, delimiter="\t" if src.lower().endswith(".tsv") else ","))}
else:
    if src.lower().endswith((".xls", ".ods")):
        out = tempfile.mkdtemp()
        subprocess.run(["soffice", f"-env:UserInstallation=file://{out}/profile", "--headless", "--convert-to", "xlsx",
                        "--outdir", out, src], capture_output=True, check=True)
        src = glob.glob(f"{out}/*.xlsx")[0]
    import openpyxl
    values, formulas = openpyxl.load_workbook(src, data_only=True), openpyxl.load_workbook(src)
    # a workbook built by code has formulas but no saved values: show the formula there
    sheets = {sheet.title: [[value if value is not None else formula for value, formula in zip(row, formula_row)]
                            for row, formula_row in zip(sheet.iter_rows(values_only=True),
                                                        formulas[sheet.title].iter_rows(values_only=True))]
              for sheet in values.worksheets}
def as_csv(rows):
    text = io.StringIO()
    csv.writer(text, lineterminator="\n").writerows(["" if cell is None else cell for cell in row] for row in rows)
    return text.getvalue()
whole = "\n".join(f"## {title}\n{as_csv(rows)}" for title, rows in sheets.items())
if len(whole.encode()) <= limit:
    print(whole)
else:
    for title, rows in sheets.items():
        width = max((len(row) for row in rows), default=0)
        print(f"## {title}\n{as_csv(rows[:20])[:limit]}[{len(rows)} rows x {width} columns; showing the first 20]\n")
    print(f"[The full data is {len(whole.encode()) // 1024} KB as CSV, over the {limit // 1024} KB limit. "
          f"It stays at {sys.argv[1]}; use bash or run_program to analyze it.]")
"""
