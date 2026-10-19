"""read's, attach's and review_file's work, done inside the sandbox where Pillow, openpyxl, LibreOffice,
poppler, pandoc and Indra live: one call in, one JSON result out, so a call is one round trip and only its
result leaves the sandbox. The host runs this file's source with `python -c` and imports it for pi's text
limits and the view hash."""
import base64
import codecs
import glob
import hashlib
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import uuid

MAX_LINES, MAX_BYTES = 2000, 50 * 1024  # pi's cut for long text
MAX_PAGES = 10  # page images per read
MAX_VIEW_PAGES = 100  # a longer file gets no view
SHEET_BYTES = 25 * 1024  # a spreadsheet up to this much CSV comes back whole
MAX_FILE = 50 * 1024 * 1024  # the most any tool pulls out of the sandbox at once
IMAGE_SIDE = 1568  # px on the long side
IMAGE_BYTES = 3_750_000  # an image passed through as is: 5 MB as base64, the strictest provider's cap
VIEWS = "/workspace/.agent/views"
SOURCE, TRANSCRIPT, DURATION, TEXT = "source.pdf", "transcript.jsonl", "duration", "text.txt"  # in a view
REVIEWS = "/workspace/.agent/reviews"  # a file per reviewed version, by hash, naming the paths it was reviewed at
REVIEWED = {"pages", "image"}  # the families deliver_file sends only once reviewed
VIEW_VERSION = b"1"  # hashed with each file's bytes: a change to what views hold builds them all anew

FAMILIES = {
    "image": {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic", ".heif", ".avif"},
    "pages": {".pdf", ".docx", ".doc", ".odt", ".rtf", ".pages", ".pptx", ".ppt", ".odp", ".key"},
    "audio": {".mp3", ".m4a", ".wav", ".aac", ".flac", ".ogg", ".opus", ".webm"},
    "book": {".epub", ".ipynb"},
    "sheet": {".xlsx", ".xlsm", ".xls", ".ods", ".numbers"},
    "email": {".eml", ".msg"},
    "video": {".mp4", ".m4v", ".mov", ".mkv", ".avi"},
    "archive": {".zip", ".tar", ".gz", ".tgz", ".bz2", ".xz", ".7z", ".rar"},
}  # anything else is text if it decodes, else binary
UNSUPPORTED = {"email": "Email files", "video": "Video files", "archive": "Archives",
               "text": "Binary files"}  # a text file that doesn't decode is binary


class Problem(Exception):
    """Something about the file the model can fix; its message goes back as the result."""


class Transcribe(Exception):
    """A recording has no view yet. The host makes it, since the sandbox has no network."""


def read(path: str, shown: str, offset: int | None, limit: int | None, pages: str | None) -> list[dict]:
    """The result's parts for the file at `path`, which the model calls `shown`."""
    _check(path, shown)
    family = family_of(path)
    if pages and family != "pages":
        raise Problem(f"{shown} has no pages. pages works on PDFs, documents and slides.")
    if family == "image":
        image = _image(_load(path))
        return [_text(f"Read image file [{image['mime']}]"), image]
    if family == "sheet":
        return [_text(_sheets(path))]
    if family in ("pages", "audio", "book"):
        view = _view(_load(path), family, shown)
        if pages:
            return _page_images(view, pages, shown)
        text, source = _transcript(view, family)
        return [_text(page(text, _shown(source), offset, limit))]
    if family == "text" and (text := _decode(_load(path))) is not None:
        return [_text(page(text, shown, offset, limit))]
    raise Problem(f"read doesn't support {UNSUPPORTED[family].lower()} yet. Use bash to work with {shown}.")


def review(path: str, shown: str) -> list[dict]:
    """review_file's parts: a PDF, document or slides as its page count, a contact sheet and each
    page's first line, or an image as itself. Marks the version it read reviewed, for deliver_file."""
    _check(path, shown)
    family, data = family_of(path), _load(path)
    if family not in REVIEWED:
        raise Problem(f"review_file checks PDFs, documents, slides and images. Use read for {shown}.")
    if family == "image":
        parts = [_text(f"Reviewed {shown}."), _image(data)]
    else:
        try:
            view = _view(data, family, shown)
        except Problem as problem:  # nothing to show, but the model has looked, and the file can still go
            parts = [_text(f"{shown} can't be shown: {problem} The review still counts: deliver_file will send it "
                           f"as it is.")]
        else:
            with open(f"{view}/{TRANSCRIPT}") as f:
                entries = [json.loads(line) for line in f]
            firsts = "\n".join(f"{entry['page']}. {_first_line(entry['text'])}" for entry in entries)
            parts = [_text(f"Reviewed {shown} ({len(entries)} pages). The first line of each page:\n{firsts}\n\n"
                           f"To see a page in detail, read it with pages."), _contact_sheet(view, len(entries))]
    os.makedirs(REVIEWS, exist_ok=True)
    with open(f"{REVIEWS}/{view_hash(data)}", "a") as f:  # the path, so deliver_file can tell changed from new
        f.write(path + "\n")
    return parts


def attach(path: str, shown: str) -> list[dict]:
    """The user message's parts for an upload: its path, and what read would show within read's limits."""
    family = family_of(path)
    if family == "image":
        return [_text(attached(shown)), _image(_load(path))]
    if family == "sheet":
        return [_text(attached(shown))]
    if family in ("pages", "audio"):
        view = _view(_load(path), family, shown)
        text, source = _transcript(view, family)
        size = f"{len(os.listdir(f'{view}/pages'))} pages" if family == "pages" else _clock(_duration(view))
        note = (f"Its transcript is too long to show here. Search {_shown(source)} with bash (rg, grep), and read "
                f"{shown} with offset and limit" + (", or with pages to see page images." if family == "pages" else "."))
        return [_text(attached(f"{shown} ({size})", text if fits(text) else note))]
    if family == "book":
        text, source = _transcript(_view(_load(path), family, shown), family)
        return [_text(attached(shown, page(text, _shown(source), None, None)))]
    if family == "text" and (text := _decode(_load(path))) is not None:
        return [_text(attached(shown, page(text, shown, None, None)))]
    return [_text(attached(shown, f"{UNSUPPORTED[family]} aren't supported yet."))]


def attached(shown: str, body: str = "") -> str:
    return f"Attached file: {shown}" + (f"\n\n{body}" if body else "")


def page(text: str, path: str, offset: int | None, limit: int | None) -> str:
    """pi's read window: from `offset`, at most `limit` lines, cut at 2000 lines or 50 KB, with a
    notice that says how to read on."""
    lines = text.split("\n")
    if text.endswith("\n"):
        lines.pop()
    total = len(lines)
    start = (offset or 1) - 1
    if offset and start >= total:
        raise Problem(f"Offset {offset} is beyond end of file ({total} lines total)")
    end = min(total, start + limit) if limit else total
    shown, size = [], 0
    for line in lines[start:end]:
        size += len(line.encode()) + 1
        if len(shown) == MAX_LINES or size > MAX_BYTES:
            break
        shown.append(line)
    if not shown and start < end:
        line = start + 1
        return (f"[Line {line} is {format_size(len(lines[start].encode()))}, exceeds {format_size(MAX_BYTES)} limit. "
                f"Use bash: sed -n '{line}p' {path} | head -c {MAX_BYTES}]")
    last = start + len(shown)
    out = "\n".join(shown)
    if last < end:
        cap = "" if len(shown) == MAX_LINES else f" ({format_size(MAX_BYTES)} limit)"
        out += (f"\n\n[Showing lines {start + 1}-{last} of {total}{cap}. Use offset={last + 1} to continue, "
                f"or search {path} with bash (rg, grep).]")
    elif end < total:
        out += f"\n\n[{total - end} more lines in file. Use offset={end + 1} to continue.]"
    return out


def fits(text: str) -> bool:
    """Whether text comes back whole, within pi's limits."""
    return text.count("\n") < MAX_LINES and len(text.encode()) <= MAX_BYTES


def format_size(n: int) -> str:
    """Sizes the way pi writes them: 512B, 50.0KB, 1.2MB."""
    if n < 1024:
        return f"{n}B"
    if n < 1024 * 1024:
        return f"{n / 1024:.1f}KB"
    return f"{n / 1024 / 1024:.1f}MB"


def family_of(path: str) -> str:
    extension = os.path.splitext(path)[1].lower()
    return next((family for family, extensions in FAMILIES.items() if extension in extensions), "text")


def view_hash(data: bytes) -> str:
    """The name of a file's view folder."""
    return hashlib.sha256(VIEW_VERSION + data).hexdigest()[:16]


def _check(path: str, shown: str) -> None:
    if not os.path.exists(path):
        raise Problem(f"File not found: {shown}")
    if os.path.isdir(path):
        raise Problem(f"{shown} is a directory. Use bash to list it.")
    size = os.path.getsize(path)
    if size > MAX_FILE:
        raise Problem(f"{shown} is {format_size(size)}, over the {format_size(MAX_FILE)} limit. Use bash to work with it.")


def _load(path: str) -> bytes:
    with open(path, "rb") as f:
        return f.read()


def _decode(data: bytes) -> str | None:
    """The bytes as text: UTF-8, UTF-16 with a BOM, or a legacy encoding charset-normalizer finds.
    None for a binary file."""
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16")
    if re.search(rb"[\x00-\x08\x0e-\x19\x1c-\x1f]", data[:8192]):  # control bytes no text holds; DOS ends at \x1a
        return None
    text = data.decode("utf-8", errors="replace").removeprefix("\ufeff")
    stray = text.count("\ufffd")
    if len(re.findall("[^\x00-\x7f\ufffd]", text)) >= stray:  # UTF-8, perhaps with a stray byte
        return text
    import charset_normalizer  # pyright: ignore[reportMissingImports]  (the sandbox has it)
    matches = charset_normalizer.from_bytes(data)
    if not (best := matches.best()):
        return text
    # a tie goes to Windows' Western code page, the commonest legacy one; charset-normalizer
    # would pick cp1250 and turn French è into č
    return str(next((match for match in matches if match.encoding == "cp1252"
                     and (match.chaos, match.coherence) == (best.chaos, best.coherence)), best))


def _view(data: bytes, family: str, shown: str) -> str:
    """The view folder of the file `shown` holding `data`, built if it's missing (TOOLS.md, Files).
    The build uses the bytes that were hashed, so a file changed meanwhile can't land under another
    version's hash. It happens in a temp folder renamed into place, so a crash leaves no half-made
    view and parallel builds don't collide."""
    view = f"{VIEWS}/{view_hash(data)}"
    if os.path.isdir(view):
        return view
    if family == "audio":
        raise Transcribe()
    work = f"{view}.{uuid.uuid4().hex[:8]}"
    os.makedirs(work)
    try:
        with tempfile.TemporaryDirectory() as local:  # the snapshot, off the volume
            snapshot = f"{local}/{os.path.basename(shown)}"
            with open(snapshot, "wb") as f:
                f.write(data)
            if family == "pages":
                _pages_view(snapshot, shown, work)
            else:  # an ebook or notebook: a plain-text copy
                _run("pandoc", snapshot, "-t", "plain", "--wrap=none", "-o", f"{work}/{TEXT}")
        try:
            os.rename(work, view)
        except OSError:  # a parallel build of the same file got there first
            if not os.path.isdir(view):
                raise
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return view


def _pages_view(path: str, shown: str, out: str) -> None:
    """source.pdf, pages/<n>.jpg and a transcript of each page's text, from Indra."""
    import indra_ocr  # pyright: ignore[reportMissingImports]
    import pypdfium2  # pyright: ignore[reportMissingImports]
    pdf = f"{out}/{SOURCE}"
    if path.lower().endswith(".pdf"):
        shutil.copyfile(path, pdf)
    else:
        with tempfile.TemporaryDirectory() as work:
            shutil.move(_convert(path, "pdf", work), pdf)
    try:
        total = len(pypdfium2.PdfDocument(pdf))
    except pypdfium2.PdfiumError as error:
        raise Problem(f"Could not open {shown}: {error}") from None
    if total > MAX_VIEW_PAGES:
        raise Problem(f"{shown} has {total} pages, too long for read, which takes {MAX_VIEW_PAGES} at most. "
                      f"Use bash to work with it.")
    os.mkdir(f"{out}/pages")
    with subprocess.Popen(["pdftoppm", "-jpeg", "-jpegopt", "quality=85", "-scale-to", str(IMAGE_SIDE), pdf,
                           f"{out}/pages/p"], stderr=subprocess.PIPE, text=True) as render:  # runs beside Indra
        try:
            [(_, pages)] = indra_ocr.convert_many([pdf])
        except BaseException:
            render.kill()
            raise
        errors = render.communicate()[1]
    if isinstance(pages, Exception):
        raise Problem(f"Could not read {shown}: {pages}")
    if render.returncode:
        raise Problem(f"pdftoppm failed: {errors.strip() or 'no reason given'}")
    with open(f"{out}/{TRANSCRIPT}", "w") as f:
        for index, result in pages:
            text = result.markdown if result.markdown is not None else result.text
            # PDFium can split an emoji into its UTF-16 halves: rejoin them, and replace a lone half
            text = text.encode("utf-16", "surrogatepass").decode("utf-16", "replace")
            f.write(json.dumps({"page": index + 1, "text": text}, ensure_ascii=False) + "\n")
    for name in os.listdir(f"{out}/pages"):  # p-01.jpg, padded to the page count's digits
        os.rename(f"{out}/pages/{name}", f"{out}/pages/{int(name.removeprefix('p-').removesuffix('.jpg'))}.jpg")


def _transcript(view: str, family: str) -> tuple[str, str]:
    """A view's text as read shows it, and the file to search for it."""
    if family == "book":
        with open(f"{view}/{TEXT}") as f:
            return f.read(), f"{view}/{TEXT}"
    with open(f"{view}/{TRANSCRIPT}") as f:
        entries = [json.loads(line) for line in f]
    if family == "pages":
        return "\n\n".join(f"## Page {entry['page']}\n{entry['text']}" for entry in entries), f"{view}/{TRANSCRIPT}"
    return "\n".join(f"[{_clock(entry['start'])}] " + (f"Speaker {entry['speaker']}: " if entry["speaker"] else "")
                     + entry["text"] for entry in entries), f"{view}/{TRANSCRIPT}"


def _duration(view: str) -> float:
    """A recording's length in seconds, as Together measured it."""
    with open(f"{view}/{DURATION}") as f:
        return float(f.read())


def _page_images(view: str, pages: str, shown: str) -> list[dict]:
    """The pages' images, each after a `Page 3 of 14` part. pages is "3" or "2-5"."""
    total = len(os.listdir(f"{view}/pages"))
    match = re.fullmatch(r"\s*(\d+)\s*(?:-\s*(\d+)\s*)?", pages)
    first, last = (int(match[1]), int(match[2] or match[1])) if match else (0, 0)
    if not 1 <= first <= last:
        raise Problem(f'Invalid pages "{pages}". Use one page like "3" or a range like "2-5"; pages start at 1.')
    if last - first >= MAX_PAGES:
        raise Problem(f'pages "{pages}" asks for {last - first + 1} pages. Ask for at most {MAX_PAGES} per call.')
    if first > total:
        raise Problem(f"{shown} has {total} pages.")
    parts = []
    for number in range(first, min(last, total) + 1):
        with open(f"{view}/pages/{number}.jpg", "rb") as f:
            parts += [_text(f"Page {number} of {total}"), _image_part(f.read(), "image/jpeg")]
    return parts


def _first_line(text: str) -> str:
    return next((line.strip(" #")[:100] for line in text.splitlines() if line.strip(" #")), "(no text)")


def _contact_sheet(view: str, total: int) -> dict:
    """Every page as a numbered thumbnail on one image, IMAGE_SIDE px at most."""
    from PIL import Image, ImageDraw, ImageFont  # pyright: ignore[reportMissingImports]
    columns = math.ceil(math.sqrt(total))
    cell = IMAGE_SIDE // columns
    sheet = Image.new("RGB", (columns * cell, math.ceil(total / columns) * cell), (200, 200, 200))
    draw, font = ImageDraw.Draw(sheet), ImageFont.load_default(size=max(12, cell // 10))
    for number in range(1, total + 1):
        x, y = (number - 1) % columns * cell, (number - 1) // columns * cell
        with Image.open(f"{view}/pages/{number}.jpg") as thumbnail:
            thumbnail.thumbnail((cell - 6, cell - 6))
            sheet.paste(thumbnail, (x + (cell - thumbnail.width) // 2, y + (cell - thumbnail.height) // 2))
        box = draw.textbbox((x + 8, y + 6), str(number), font=font)
        draw.rectangle((box[0] - 4, box[1] - 3, box[2] + 4, box[3] + 3), fill="black")
        draw.text((x + 8, y + 6), str(number), fill="white", font=font)
    out = io.BytesIO()
    sheet.save(out, "JPEG", quality=85)
    return _image_part(out.getvalue(), "image/jpeg")


def _clock(seconds: float) -> str:
    minutes, second = divmod(int(seconds), 60)
    return f"{minutes // 60:02d}:{minutes % 60:02d}:{second:02d}"


def _shown(path: str) -> str:
    return path.removeprefix("/workspace/")


def _image(data: bytes) -> dict:
    """The image part, under the shared image rule."""
    from PIL import Image, ImageOps  # pyright: ignore[reportMissingImports]  (the sandbox has these)
    import pillow_heif  # pyright: ignore[reportMissingImports]
    pillow_heif.register_heif_opener()
    image = Image.open(io.BytesIO(data))
    upright = image.getexif().get(0x0112, 1) == 1  # models read no EXIF, so a turned photo must be turned
    if (image.format in ("PNG", "JPEG", "GIF", "WEBP") and max(image.size) <= IMAGE_SIDE and upright
            and len(data) <= IMAGE_BYTES):
        return _image_part(data, Image.MIME[image.format])
    # too big, turned, or a format providers don't take: redraw it, JPEG unless it's transparent
    image = ImageOps.exif_transpose(image)
    image.thumbnail((IMAGE_SIDE, IMAGE_SIDE))
    data, mime = b"", "image/png"
    if image.mode in ("RGBA", "LA") or "transparency" in image.info:
        image = image.convert("RGBA")
        data = _encoded(image, "PNG")
    if not data or len(data) > IMAGE_BYTES:  # a transparent photo can top 9 MB as PNG: JPEG it on white
        if image.mode == "RGBA":  # dropping alpha alone would show whatever color the clear areas hide
            image = Image.alpha_composite(Image.new("RGBA", image.size, "white"), image)
        data, mime = _encoded(image.convert("RGB"), "JPEG"), "image/jpeg"
    return _image_part(data, mime)


def _encoded(image, format: str) -> bytes:
    out = io.BytesIO()
    image.save(out, format, quality=85)
    return out.getvalue()


def _sheets(path: str) -> str:
    """Every sheet as CSV when it all fits in 25 KB, else a preview of each, cut at read's limits."""
    import csv
    import warnings
    import openpyxl  # pyright: ignore[reportMissingModuleSource]
    warnings.simplefilter("ignore")  # openpyxl warns about every extension it skips
    with tempfile.TemporaryDirectory() as work:
        book = path if path.lower().endswith((".xlsx", ".xlsm")) else _convert(path, "xlsx", work)
        values, formulas = openpyxl.load_workbook(book, data_only=True), openpyxl.load_workbook(book)
    # a workbook built by code has formulas but no saved values: show the formula there
    sheets = {sheet.title: [[value if value is not None else formula for value, formula in zip(row, formula_row)]
                            for row, formula_row in zip(sheet.iter_rows(values_only=True),
                                                        formulas[sheet.title].iter_rows(values_only=True))]
              for sheet in values.worksheets}

    def as_csv(rows: list) -> str:
        text = io.StringIO()
        csv.writer(text, lineterminator="\n").writerows(["" if cell is None else cell for cell in row] for row in rows)
        return text.getvalue()

    whole = "\n".join(f"## {title}\n{as_csv(rows)}" for title, rows in sheets.items()).rstrip("\n")
    if len(whole.encode()) <= SHEET_BYTES:
        return whole
    previews = "\n\n".join(f"## {title}\n{as_csv(rows[:20])}[{len(rows)} rows x {max(map(len, rows), default=0)} "
                           f"columns; showing the first 20]" for title, rows in sheets.items())
    notice = (f"\n\n[The full data is {len(whole.encode()) // 1024} KB as CSV, over the {SHEET_BYTES // 1024} KB "
              f"limit. It stays at {path}; use bash or run_program to analyze it.]")
    lines = previews.split("\n")[:MAX_LINES - 2]  # the notice takes the last two lines
    return "\n".join(lines).encode()[:MAX_BYTES - len(notice.encode())].decode(errors="ignore") + notice


def _convert(path: str, to: str, work: str) -> str:
    """LibreOffice's copy of the file in another format, with a profile of its own so parallel
    conversions don't fight over one."""
    _run("soffice", f"-env:UserInstallation=file://{work}/profile", "--headless", "--convert-to", to,
         "--outdir", work, path)
    converted = glob.glob(f"{work}/*.{to}")
    if not converted:
        raise Problem(f"LibreOffice couldn't open {os.path.basename(path)}.")
    return converted[0]


def _run(*args: str) -> None:
    done = subprocess.run(args, capture_output=True, text=True)
    if done.returncode:
        raise Problem(f"{args[0]} failed: {done.stderr.strip() or 'no reason given'}")


def _text(text: str) -> dict:
    return {"type": "text", "text": text}


def _image_part(data: bytes, mime: str) -> dict:
    return {"type": "image", "data": base64.b64encode(data).decode(), "mime": mime}


if __name__ == "__main__":
    request = json.loads(sys.argv[1])
    try:
        print(json.dumps({"content": {"read": read, "attach": attach, "review": review}[request.pop("do")](**request)}))
    except Problem as problem:
        print(json.dumps({"problem": str(problem)}))
    except Transcribe:
        print(json.dumps({"transcribe": True}))
