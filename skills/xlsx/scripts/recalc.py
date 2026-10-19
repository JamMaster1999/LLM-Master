"""Computes every formula in a workbook with LibreOffice and saves the values into the file, then
lists the cells that came out as errors. openpyxl writes formulas without values, so until this
runs, previews and pandas see empty cells.

    python recalc.py gradebook.xlsx

The file is rewritten in place, with its formatting, charts and validation kept. Exits 1 when a
formula comes out as an error."""
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

SHOWN = 30  # error cells listed; the count covers them all


def main(path: str) -> int:
    if not path.lower().endswith(".xlsx"):
        sys.exit(f"{path} isn't an .xlsx file. Convert it first, which keeps a copy of the original: "
                 "soffice --headless --convert-to xlsx --outdir . FILE")
    if not os.path.isfile(path):
        sys.exit(f"{path} doesn't exist.")
    with tempfile.TemporaryDirectory() as work:
        # its own profile, so a stuck LibreOffice elsewhere can't block this one
        done = subprocess.run(["soffice", f"-env:UserInstallation=file://{work}/profile", "--headless", "--calc",
                               "--convert-to", "xlsx", "--outdir", work, path], capture_output=True, text=True,
                              timeout=300)
        result = os.path.join(work, os.path.splitext(os.path.basename(path))[0] + ".xlsx")  # .XLSX comes back lowercase
        if not os.path.isfile(result):
            sys.exit(f"LibreOffice couldn't open {path}: {(done.stderr or done.stdout).strip()}")
        shutil.move(result, path)
    formulas, values = load_workbook(path), load_workbook(path, data_only=True)
    count, errors = 0, []
    for sheet in formulas.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":  # array formulas too
                    count += 1
                    value = values[sheet.title][cell.coordinate]
                    if value.data_type == "e":
                        formula = getattr(cell.value, "text", cell.value)
                        errors.append(f"{sheet.title}!{cell.coordinate} {value.value}: {formula}")
    print(f"Recalculated {count} formulas in {path}.")
    if not errors:
        print("No errors.")
        return 0
    print(f"{len(errors)} formulas give errors:")
    print("\n".join(errors[:SHOWN]))
    if len(errors) > SHOWN:
        print(f"... and {len(errors) - SHOWN} more.")
    print("A formula LibreOffice couldn't read comes back in lowercase: an unknown sheet or function, or a newer "
          "function missing its _xlfn. prefix.")
    return 1


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    sys.exit(main(sys.argv[1]))
